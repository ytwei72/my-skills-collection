# -*- coding: utf-8 -*-
"""Markdown -> 规范、美观、可读的 Word(.docx)。

用法:
    python md_to_docx.py <input.md> [output.docx]

版式（与「应急包挂硬件方案」一版对齐）:
- 微软雅黑；藏青 / 青绿 / 橙色。
- 导航窗格：章、节必须是正文中的「标题 1 / 标题 2」段落，禁止放进表格。
- 仅有一个 #、且有多个 ## 时，# 作封面标题（不进大纲），## 为一级，### 为二级。
- 表头深色、斑马纹、列宽按内容自适应；链接可点击；页眉细线、页脚页码。

依赖: python-docx。图片按比例缩放时若已安装 Pillow 会使用，否则按固定宽度嵌入。
"""
from __future__ import annotations

import os
import re
import sys
from dataclasses import dataclass, field

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

try:
    from PIL import Image
except ImportError:
    Image = None

# 配色与交通公路剧本 / 教材排版版一致
NAVY = "1B3A4B"
TEAL = "2A6F7F"
ORANGE = "C45C26"
BODY = "2C3338"
MUTED = "5B6770"
WHITE = "FFFFFF"
FILL_ICE = "E8F1F4"
FILL_SAND = "F7F4EE"
FILL_MINT = "EEF6F3"
FILL_ROW = "F4F7F8"
LINE = "C5D5DC"
FONT = "微软雅黑"
FONT_MONO = "Consolas"

COLOR_NAVY = RGBColor.from_string(NAVY)
COLOR_TEAL = RGBColor.from_string(TEAL)
COLOR_BODY = RGBColor.from_string(BODY)
COLOR_MUTED = RGBColor.from_string(MUTED)
COLOR_WHITE = RGBColor.from_string(WHITE)

PAGE_W = 21.0
ML = MR = 1.8
MT = 2.0
MB = 1.7

HIGHLIGHT = {"判定", "对本方案", "相对按键型", "结论", "合计", "按键型，按零售模组"}
LABEL_HEADERS = {"字段", "项目", "项", "码", "块", "组成", "物料", "物料或整机", "加装", "路径"}
LINK_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
TOKEN = re.compile(
    r"\*\*(.+?)\*\*"
    r"|`([^`]+)`"
    r"|\[([^\]]+)\]\(([^)]+)\)"
    r"|(https?://[^\s<>\[\]（）()，。；、]+)"
    r"|(?<!\*)\*([^*]+)\*(?!\*)"
)
IMG_RE = re.compile(
    r'(?:!\[([^\]]*)\]\(([^)]+)\)|<img\s+[^>]*src="([^"]+)"[^>]*alt="([^"]*)"[^>]*>)',
    re.I,
)
_bookmark_seq = 0


def cm_twips(cm: float) -> int:
    return int(round(cm * 1440 / 2.54))


def display_width(text: str) -> int:
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in text)


def plain(text: str) -> str:
    t = re.sub(r"!\[[^\]]*\]\([^)]+\)", "", text)
    t = re.sub(r"<img\b[^>]*>", "", t, flags=re.I)
    t = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", t)
    t = t.replace("**", "").replace("`", "")
    t = re.sub(r"(?<!\*)\*(?!\*)", "", t)
    return t.strip()


def set_run_font(run, size, color, *, bold=False, italic=False, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = color if isinstance(color, RGBColor) else RGBColor.from_string(color)
    run.bold = bold or None
    run.italic = italic or None
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), FONT if name == FONT_MONO else name)
    rfonts.set(qn("w:cs"), name)


def _rpr(size, color_hex, *, bold=False, italic=False, underline=False, name=FONT):
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    east = FONT if name == FONT_MONO else name
    for attr, val in (("w:ascii", name), ("w:hAnsi", name), ("w:eastAsia", east), ("w:cs", name)):
        rfonts.set(qn(attr), val)
    rpr.append(rfonts)
    for tag in ("w:sz", "w:szCs"):
        el = OxmlElement(tag)
        el.set(qn("w:val"), str(int(round(size * 2))))
        rpr.append(el)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color_hex)
    rpr.append(c)
    if bold:
        rpr.append(OxmlElement("w:b"))
        rpr.append(OxmlElement("w:bCs"))
    if italic:
        rpr.append(OxmlElement("w:i"))
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
    return rpr


def _rel_id(part, url: str) -> str:
    cache = getattr(part, "_href_cache", None)
    if cache is None:
        cache = {}
        part._href_cache = cache
    if url not in cache:
        cache[url] = part.relate_to(url, LINK_REL, is_external=True)
    return cache[url]


def add_hyperlink(paragraph, text, url, *, size, bold=False, italic=False):
    link = OxmlElement("w:hyperlink")
    if url.startswith(("http://", "https://")):
        link.set(qn("r:id"), _rel_id(paragraph.part, url))
    else:
        link.set(qn("w:anchor"), url)
    run = OxmlElement("w:r")
    run.append(_rpr(size, TEAL, bold=bold, italic=italic, underline=url.startswith("http")))
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, text, *, size=10.5, color=BODY, bold=False, italic=False, emph=NAVY):
    if not text:
        return
    pos = 0
    base = color if isinstance(color, str) else str(color)
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            set_run_font(run, size, base, bold=bold, italic=italic)
        if m.group(1) is not None:
            run = paragraph.add_run(m.group(1))
            set_run_font(run, size, emph, bold=True, italic=italic)
        elif m.group(2) is not None:
            run = paragraph.add_run(m.group(2))
            set_run_font(run, size - 0.5, ORANGE, name=FONT_MONO)
        elif m.group(3) is not None:
            add_hyperlink(paragraph, m.group(3), m.group(4), size=size, bold=bold, italic=italic)
        elif m.group(5) is not None:
            add_hyperlink(paragraph, m.group(5), m.group(5), size=size, bold=bold, italic=italic)
        elif m.group(6) is not None:
            run = paragraph.add_run(m.group(6))
            set_run_font(run, size, MUTED, italic=True)
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size, base, bold=bold, italic=italic)


def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcpr.append(mar)


def nil_table_borders(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    table._tbl.tblPr.append(borders)


def set_table_borders(table, color=LINE, sz="4"):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    table._tbl.tblPr.append(borders)


def set_col_widths(table, widths_cm):
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(cm_twips(sum(widths_cm))))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    jc = tbl_pr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        tbl_pr.append(jc)
    jc.set(qn("w:val"), "center")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(1, grid)
    for child in list(grid):
        grid.remove(child)
    for w in widths_cm:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(cm_twips(w)))
        grid.append(gc)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(cm_twips(widths_cm[i])))
            tc_w.set(qn("w:type"), "dxa")


def row_flags(row, *, header=False):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))
    if header:
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        tr_pr.append(th)


def para_border(paragraph, edge, color, sz="12", space="1"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), sz)
    el.set(qn("w:space"), space)
    el.set(qn("w:color"), color)
    p_bdr.append(el)


def para_shade(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def bookmark(paragraph, name: str):
    global _bookmark_seq
    _bookmark_seq += 1
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(_bookmark_seq))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(_bookmark_seq))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def style_font(style, size, color_hex, *, bold=True):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor.from_string(color_hex)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)


def ensure_outline(style, level_index: int):
    p_pr = style.element.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = OxmlElement("w:pPr")
        style.element.append(p_pr)
    for old in p_pr.findall(qn("w:outlineLvl")):
        p_pr.remove(old)
    ol = OxmlElement("w:outlineLvl")
    ol.set(qn("w:val"), str(level_index))
    p_pr.append(ol)


def setup_heading_styles(doc: Document):
    """标题样式写在样式定义上。渲染时必须落在正文段落，不能放进表格。"""
    specs = [
        ("Heading 1", 13.5, NAVY, 0),
        ("Heading 2", 12, TEAL, 1),
        ("Heading 3", 11, TEAL, 2),
        ("Heading 4", 10.5, TEAL, 3),
    ]
    for name, size, color, level in specs:
        style = doc.styles[name]
        style_font(style, size, color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.15
        ensure_outline(style, level)


def content_width(section) -> float:
    return (section.page_width - section.left_margin - section.right_margin) / 360000.0


def compute_col_widths(rows, total_cm, *, image=False, min_cm=2.15, cap_units=32):
    ncol = max(len(r) for r in rows)
    if image or ncol == 0:
        return [total_cm / max(ncol, 1)] * ncol
    nat = []
    for c in range(ncol):
        longest = 0
        for r in rows:
            if c < len(r):
                longest = max(longest, display_width(plain(r[c])))
        nat.append(max(4, min(longest, cap_units)))
    if nat[0] <= 8:
        nat[0] = max(nat[0], 6)
    raw = [max(min_cm, total_cm * n / sum(nat)) for n in nat]
    scale = total_cm / sum(raw)
    return [w * scale for w in raw]


def fit_image(path, max_w, max_h):
    if Image is None or not os.path.exists(path):
        return max_w
    with Image.open(path) as im:
        w, h = im.size
    if h <= 0:
        return max_w
    width = max_w
    height = width * h / w
    if height > max_h:
        width = max_h * w / h
    return width


def is_image_cell(text: str) -> bool:
    s = text.strip()
    return s.startswith("<img") or s.startswith("![")


@dataclass
class Block:
    kind: str
    text: str = ""
    rows: list = field(default_factory=list)
    items: list = field(default_factory=list)
    bookmark: str = ""
    level: int = 0


def split_row(line: str):
    parts = line.strip()
    if parts.startswith("|"):
        parts = parts[1:]
    if parts.endswith("|"):
        parts = parts[:-1]
    return [c.strip() for c in parts.split("|")]


def is_sep(cells) -> bool:
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells)


def parse(md: str) -> list[Block]:
    lines = md.replace("\r\n", "\n").split("\n")
    blocks: list[Block] = []
    i, n = 0, len(lines)
    while i < n:
        s = lines[i].strip()
        if not s or s in ("---", "***"):
            i += 1
            continue
        if s.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append(Block("code", "\n".join(buf)))
            continue
        if s.startswith("|"):
            raw = []
            while i < n and lines[i].strip().startswith("|"):
                cells = split_row(lines[i])
                if not is_sep(cells):
                    raw.append(cells)
                i += 1
            if raw:
                width = max(len(r) for r in raw)
                raw = [r + [""] * (width - len(r)) for r in raw]
                blocks.append(Block("table", rows=raw))
            continue
        if s.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(re.sub(r"^>\s?", "", lines[i].strip()))
                i += 1
            blocks.append(Block("quote", items=buf))
            continue
        if s.startswith("#### "):
            blocks.append(Block("h", s[5:].strip(), level=4))
            i += 1
            continue
        if s.startswith("### "):
            blocks.append(Block("h", s[4:].strip(), level=3))
            i += 1
            continue
        if s.startswith("## "):
            blocks.append(Block("h", s[3:].strip(), level=2))
            i += 1
            continue
        if s.startswith("# "):
            blocks.append(Block("h", s[2:].strip(), level=1))
            i += 1
            continue
        if re.match(r"^[-*+]\s+", s):
            items = []
            while i < n and re.match(r"^[-*+]\s+", lines[i].strip()):
                items.append(re.sub(r"^[-*+]\s+", "", lines[i].strip()))
                i += 1
            blocks.append(Block("ul", items=items))
            continue
        if re.match(r"^\d+\.\s", s):
            items = []
            while i < n and re.match(r"^\d+\.\s", lines[i].strip()):
                items.append(re.sub(r"^\d+\.\s*", "", lines[i].strip()))
                i += 1
            blocks.append(Block("ol", items=items))
            continue
        img = IMG_RE.search(s)
        if img and (s.startswith("![") or s.lower().startswith("<img")):
            alt = img.group(1) if img.group(1) is not None else (img.group(4) or "")
            src = img.group(2) if img.group(2) is not None else img.group(3)
            blocks.append(Block("image", alt or "", rows=[[src]]))
            i += 1
            continue
        if s.startswith("*") and s.endswith("*") and not s.startswith("**"):
            blocks.append(Block("caption", s[1:-1].strip()))
            i += 1
            continue
        blocks.append(Block("p", s))
        i += 1
    return blocks


def report_mode(blocks: list[Block]) -> bool:
    """方案/调研稿通常只有一个文档标题，章用 ##。此时 ## 才是导航里的一级。"""
    h1 = sum(1 for b in blocks if b.kind == "h" and b.level == 1)
    h2 = sum(1 for b in blocks if b.kind == "h" and b.level == 2)
    return h1 == 1 and h2 >= 2


def outline_level(markdown_level: int, report: bool) -> int | None:
    """返回 Word 大纲级别（1 起）。None 表示不进入导航。"""
    if report:
        if markdown_level == 1:
            return None
        return markdown_level - 1  # ## -> 1, ### -> 2, #### -> 3
    return markdown_level  # # -> 1, ## -> 2


def add_body_paragraph(doc, text="", *, size=10.5, color=BODY, before=0, after=6, align="left", italic=False, keep_next=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.28
    pf.widow_control = True
    if keep_next:
        pf.keep_with_next = True
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        add_inline(p, text, size=size, color=color, italic=italic)
    return p


def render_outline_heading(doc, text: str, word_level: int, bookmark_name: str, *, title=False):
    style = f"Heading {min(word_level, 4)}"
    p = doc.add_paragraph(text, style=style)
    pf = p.paragraph_format
    pf.keep_with_next = True
    pf.line_spacing = 1.15
    if title:
        pf.space_before = Pt(8)
        pf.space_after = Pt(10)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            set_run_font(run, 22, NAVY, bold=True)
        para_border(p, "bottom", NAVY, sz="12", space="4")
    elif word_level == 1:
        pf.space_before = Pt(14)
        pf.space_after = Pt(8)
        pf.left_indent = Cm(0.12)
        for run in p.runs:
            set_run_font(run, 13.5, NAVY, bold=True)
        para_shade(p, FILL_ICE)
        para_border(p, "left", ORANGE, sz="24", space="8")
    else:
        pf.space_before = Pt(12 if word_level == 2 else 8)
        pf.space_after = Pt(4)
        for run in p.runs:
            set_run_font(run, 12 if word_level == 2 else 11, TEAL, bold=True)
        if word_level == 2:
            para_border(p, "bottom", TEAL, sz="8", space="1")
    bookmark(p, bookmark_name)
    return p


def render_cover(doc, title: str, quote_lines: list[str]):
    table = doc.add_table(rows=2, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    nil_table_borders(table)
    width = PAGE_W - ML - MR
    set_col_widths(table, [width])
    top, bar = table.cell(0, 0), table.cell(1, 0)
    shade_cell(top, NAVY)
    shade_cell(bar, ORANGE)
    for cell, fill, h in ((top, NAVY, "200"), (bar, ORANGE, "70")):
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(" ")
        set_run_font(run, 2, fill)
        tr_pr = cell._tc.getparent().get_or_add_trPr()
        tr_h = OxmlElement("w:trHeight")
        tr_h.set(qn("w:val"), h)
        tr_h.set(qn("w:hRule"), "exact")
        tr_pr.append(tr_h)
    add_body_paragraph(doc, title, size=22, color=NAVY, before=12, after=4, align="center", keep_next=True)
    for run in doc.paragraphs[-1].runs:
        run.bold = True
    if quote_lines:
        pairs = []
        for line in quote_lines:
            if "：" in line:
                k, v = line.split("：", 1)
                pairs.append((k if len(k) <= 8 else "说明", v if len(k) <= 8 else line))
            elif line.startswith("价格"):
                pairs.append(("价格", line))
            else:
                pairs.append(("说明", line))
        kv(doc, pairs)


def kv(doc, pairs, label_w=2.8):
    width = PAGE_W - ML - MR
    table = doc.add_table(rows=len(pairs), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    nil_table_borders(table)
    set_col_widths(table, [label_w, width - label_w])
    for i, (k, v) in enumerate(pairs):
        c0, c1 = table.rows[i].cells
        shade_cell(c0, TEAL)
        shade_cell(c1, FILL_ICE if i % 2 == 0 else WHITE)
        set_cell_margins(c0)
        set_cell_margins(c1, left=100, right=100)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p0.add_run(k)
        set_run_font(run, 10, WHITE, bold=True)
        add_inline(c1.paragraphs[0], v, size=10.5, color=BODY)
        row_flags(table.rows[i])
    add_body_paragraph(doc, "", size=4, after=6)


def render_toc(doc, heads: list[tuple[str, str]]):
    p = add_body_paragraph(doc, "目录", size=12, color=TEAL, before=4, after=4, keep_next=True)
    for run in p.runs:
        run.bold = True
    para_border(p, "bottom", TEAL, sz="8", space="1")
    n = len(heads)
    rows_n = (n + 1) // 2
    width = PAGE_W - ML - MR
    table = doc.add_table(rows=rows_n, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    nil_table_borders(table)
    set_col_widths(table, [width / 2, width / 2])
    for i, (title, anchor) in enumerate(heads):
        cell = table.rows[i // 2].cells[i % 2]
        shade_cell(cell, FILL_ICE if (i // 2) % 2 == 0 else WHITE)
        set_cell_margins(cell, top=50, bottom=50, left=90, right=80)
        add_hyperlink(cell.paragraphs[0], title, anchor, size=10.5)
        row_flags(table.rows[i // 2])
    add_body_paragraph(doc, "", size=4, after=6)


def write_cell(cell, text, *, size, color, bold=False, align="left", emph=NAVY, base_dir=""):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if is_image_cell(text):
        img = IMG_RE.search(text)
        alt = ""
        src = ""
        if img:
            alt = img.group(1) if img.group(1) is not None else (img.group(4) or "")
            src = img.group(2) if img.group(2) is not None else (img.group(3) or "")
        path = src if os.path.isabs(src) else os.path.join(base_dir, src)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if os.path.exists(path):
            try:
                col_cm = 4.8
                tc_w = cell._tc.tcPr.find(qn("w:tcW")) if cell._tc.tcPr is not None else None
                if tc_w is not None and tc_w.get(qn("w:type")) == "dxa":
                    col_cm = max(2.4, int(tc_w.get(qn("w:w"))) / 567 - 0.4)
                w = fit_image(path, min(col_cm, 5.2), 4.6)
                p.add_run().add_picture(path, width=Cm(w))
            except Exception:
                add_inline(p, alt or src, size=size, color=ORANGE, bold=True)
        else:
            add_inline(p, alt or f"[缺失图片: {src}]", size=size, color=ORANGE)
        if alt:
            cap = cell.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cap.add_run(alt)
            set_run_font(run, 8, MUTED, italic=True)
        return
    add_inline(p, text, size=size, color=color, bold=bold, emph=emph)


def render_table(doc, rows, total_cm, base_dir):
    header, body = rows[0], rows[1:]
    n = len(header)
    image = any(is_image_cell(c) for r in body for c in r)
    widths = compute_col_widths(rows, total_cm, image=image, min_cm=2.0 if n >= 5 else 2.3)
    size = 8.5 if n >= 5 else (9.5 if n == 4 else 10)
    pad_v, pad_h = (46, 56) if n >= 5 else (60, 74)
    emphasize_first = (not image) and (
        plain(header[0]) in LABEL_HEADERS
        or all(display_width(plain(r[0])) <= 18 for r in body)
    )
    center_first = plain(header[0]) in {"码"} or all(display_width(plain(r[0])) <= 4 for r in body)
    table = doc.add_table(rows=1 + len(body), cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_col_widths(table, widths)
    set_table_borders(table)
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        shade_cell(cell, NAVY)
        set_cell_margins(cell, top=pad_v, bottom=pad_v, left=pad_h, right=pad_h)
        write_cell(cell, text, size=size, color=WHITE, bold=True, align="center", emph=WHITE, base_dir=base_dir)
    row_flags(table.rows[0], header=True)
    for ri, row in enumerate(body):
        label = plain(row[0]) if row else ""
        hot = label in HIGHLIGHT
        for ci in range(n):
            text = row[ci] if ci < len(row) else ""
            cell = table.rows[ri + 1].cells[ci]
            if hot:
                fill = FILL_SAND
            elif image:
                fill = WHITE
            else:
                fill = FILL_ICE if ri % 2 == 0 else WHITE
            shade_cell(cell, fill)
            set_cell_margins(cell, top=pad_v, bottom=pad_v, left=pad_h, right=pad_h)
            if ci == 0 and emphasize_first and not is_image_cell(text):
                write_cell(
                    cell, text, size=size, color=TEAL, bold=True,
                    align="center" if center_first else "left", emph=TEAL, base_dir=base_dir,
                )
            else:
                write_cell(cell, text, size=size, color=BODY, base_dir=base_dir)
        row_flags(table.rows[ri + 1])
    add_body_paragraph(doc, "", size=3, after=8)


def render_steps(doc, items, total_cm):
    table = doc.add_table(rows=len(items), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    nil_table_borders(table)
    set_col_widths(table, [1.15, total_cm - 1.15])
    for i, item in enumerate(items):
        c0, c1 = table.rows[i].cells
        shade_cell(c0, TEAL)
        shade_cell(c1, FILL_ICE if i % 2 == 0 else WHITE)
        set_cell_margins(c0, top=50, bottom=50, left=40, right=40)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=90)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p0.add_run(str(i + 1))
        set_run_font(run, 11, WHITE, bold=True)
        c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        add_inline(c1.paragraphs[0], item, size=10.5, color=BODY)
        row_flags(table.rows[i])
    add_body_paragraph(doc, "", size=3, after=8)


def render_bullets(doc, items):
    for i, item in enumerate(items):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(1)
        pf.space_after = Pt(8 if i == len(items) - 1 else 2)
        pf.line_spacing = 1.25
        pf.left_indent = Cm(0.55)
        pf.first_line_indent = Cm(-0.35)
        mark = p.add_run("●  ")
        set_run_font(mark, 8, TEAL, bold=True)
        size, color = (9.5, MUTED) if item.startswith("来源") else (10.5, BODY)
        add_inline(p, item, size=size, color=color)


def render_quote(doc, lines, total_cm):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    nil_table_borders(table)
    set_col_widths(table, [0.18, total_cm - 0.18])
    c0, c1 = table.cell(0, 0), table.cell(0, 1)
    shade_cell(c0, TEAL)
    shade_cell(c1, FILL_MINT)
    set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
    first = True
    for line in lines:
        p = c1.paragraphs[0] if first else c1.add_paragraph()
        first = False
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.28
        add_inline(p, line, size=10.5, color=NAVY)
    row_flags(table.rows[0])
    add_body_paragraph(doc, "", size=3, after=8)


def render_code(doc, text, total_cm):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    nil_table_borders(table)
    set_col_widths(table, [0.18, total_cm - 0.18])
    c0, c1 = table.cell(0, 0), table.cell(0, 1)
    shade_cell(c0, ORANGE)
    shade_cell(c1, FILL_SAND)
    set_cell_margins(c1, top=70, bottom=70, left=120, right=120)
    p = c1.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    for idx, line in enumerate(text.split("\n")):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line if line else " ")
        set_run_font(run, 9, BODY, name=FONT_MONO)
    row_flags(table.rows[0])
    add_body_paragraph(doc, "", size=3, after=8)


def render_image(doc, src, alt, base_dir, total_cm):
    path = src if os.path.isabs(src) else os.path.join(base_dir, src)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    if os.path.exists(path):
        try:
            w = fit_image(path, min(total_cm - 1, 12.2), 5.6)
            p.add_run().add_picture(path, width=Cm(w))
        except Exception:
            add_inline(p, f"[图片无法嵌入: {src}]", size=10.5, color=MUTED)
    else:
        add_inline(p, alt or f"[缺失图片: {src}]", size=10.5, color=ORANGE)
    if alt:
        add_body_paragraph(doc, alt, size=9, color=MUTED, italic=True, align="center", before=1, after=8)


def header_footer(doc, title: str, *, cover: bool):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = cover
    sec.header_distance = Cm(0.7)
    sec.footer_distance = Cm(0.55)
    header = sec.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(PAGE_W - ML - MR), WD_TAB_ALIGNMENT.RIGHT)
    hp.paragraph_format.space_after = Pt(2)
    r1 = hp.add_run(title[:40])
    set_run_font(r1, 9, TEAL, bold=True)
    r2 = hp.add_run("\t")
    set_run_font(r2, 9, MUTED)
    para_border(hp, "bottom", TEAL, sz="8", space="1")
    if cover:
        sec.first_page_header.is_linked_to_previous = False

    def footer(ft):
        fp = ft.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para_border(fp, "top", LINE, sz="6", space="1")
        lead = fp.add_run("— ")
        set_run_font(lead, 9, MUTED)
        for kind, text in (("begin", None), ("instr", " PAGE "), ("separate", None), ("show", "1"), ("end", None)):
            run = fp.add_run(text or "")
            set_run_font(run, 9, MUTED)
            if kind == "instr":
                fld = OxmlElement("w:instrText")
                fld.set(qn("xml:space"), "preserve")
                fld.text = " PAGE "
                run._r.append(fld)
            elif kind == "show":
                continue
            elif kind == "separate":
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), "separate")
                run._r.append(fc)
            else:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), kind)
                run._r.append(fc)
        tail = fp.add_run(" —")
        set_run_font(tail, 9, MUTED)

    footer(sec.footer)
    if cover:
        footer(sec.first_page_footer)


def build(src_path, out_path):
    with open(src_path, encoding="utf-8") as f:
        blocks = parse(f.read())
    base_dir = os.path.dirname(os.path.abspath(src_path))
    report = report_mode(blocks)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(PAGE_W)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(MT)
    sec.bottom_margin = Cm(MB)
    sec.left_margin = Cm(ML)
    sec.right_margin = Cm(MR)
    total = content_width(sec)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = COLOR_BODY
    rfonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.2
    setup_heading_styles(doc)
    try:
        hyper = doc.styles["Hyperlink"]
        hyper.font.color.rgb = COLOR_TEAL
        hyper.font.name = FONT
    except KeyError:
        pass

    title = os.path.splitext(os.path.basename(src_path))[0]
    for b in blocks:
        if b.kind == "h" and b.level == 1:
            title = b.text
            break

    chapters = []
    seq = 0
    for b in blocks:
        if b.kind != "h":
            continue
        word_level = outline_level(b.level, report)
        if word_level == 1:
            seq += 1
            b.bookmark = f"sec{seq}"
            chapters.append((b.text, b.bookmark))
        elif word_level:
            seq += 1
            b.bookmark = f"h{seq}"

    header_footer(doc, title, cover=report)
    if report:
        quote = []
        rest = []
        seen = False
        for b in blocks:
            if b.kind == "h" and b.level == 1:
                seen = True
                continue
            if not seen and b.kind == "quote" and not quote:
                quote = b.items
                continue
            rest.append(b)
        render_cover(doc, title, quote)
        if len(chapters) >= 3:
            render_toc(doc, chapters)
        blocks = rest

    for b in blocks:
        if b.kind == "h":
            word_level = outline_level(b.level, report)
            if word_level is None:
                add_body_paragraph(doc, b.text, size=22, color=NAVY, before=6, after=8, align="center", keep_next=True)
                for run in doc.paragraphs[-1].runs:
                    run.bold = True
            else:
                render_outline_heading(
                    doc, b.text, word_level, b.bookmark or f"h{id(b)}",
                    title=(not report and word_level == 1 and b.level == 1),
                )
        elif b.kind == "p":
            if b.text.startswith("来源"):
                add_body_paragraph(doc, b.text, size=9.5, color=MUTED, before=2, after=6)
            else:
                add_body_paragraph(doc, b.text)
        elif b.kind == "table":
            render_table(doc, b.rows, total, base_dir)
        elif b.kind == "ul":
            render_bullets(doc, b.items)
        elif b.kind == "ol":
            render_steps(doc, b.items, total)
        elif b.kind == "quote":
            render_quote(doc, b.items, total)
        elif b.kind == "code":
            render_code(doc, b.text, total)
        elif b.kind == "image":
            render_image(doc, b.rows[0][0], b.text, base_dir, total)
        elif b.kind == "caption":
            add_body_paragraph(doc, b.text, size=9, color=MUTED, italic=True, align="center", before=2, after=8)

    doc.core_properties.title = title
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    doc.save(out_path)
    return out_path


def main(argv):
    if len(argv) < 2:
        print(__doc__)
        return 1
    src = argv[1]
    if not os.path.exists(src):
        print(f"输入文件不存在: {src}")
        return 1
    out = argv[2] if len(argv) > 2 else os.path.splitext(src)[0] + ".docx"
    build(src, out)
    print("saved:", out)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
